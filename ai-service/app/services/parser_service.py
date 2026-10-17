import io
import re
from typing import List, Dict, Any, Optional
import pypdf
import docx

from app.schemas.resume import (
    ResumeParseResponse,
    ExtractedEducation,
    ExtractedProject,
    ExtractedExperience,
)
from app.services.matcher_service import normalize_token

# Known technical and soft skills taxonomy for precise extraction
KNOWN_SKILLS = [
    "python", "javascript", "typescript", "java", "c++", "c", "c#", "go", "ruby", "php", "rust",
    "react", "react.js", "angular", "vue.js", "next.js", "node.js", "express.js", "django", "flask",
    "fastapi", "spring boot", "asp.net", "html", "css", "tailwind", "bootstrap", "sass",
    "sql", "mysql", "postgresql", "mongodb", "redis", "cassandra", "sqlite", "oracle",
    "docker", "kubernetes", "aws", "azure", "gcp", "git", "github", "gitlab", "ci/cd", "jenkins", "linux",
    "machine learning", "deep learning", "nlp", "computer vision", "tensorflow", "pytorch", "scikit-learn",
    "pandas", "numpy", "power bi", "tableau", "excel", "data analysis",
    "communication", "leadership", "problem solving", "teamwork", "agile", "scrum", "rest api", "graphql"
]

EMAIL_REGEX = r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+"
PHONE_REGEX = r"(?:\+91[\-\s]?)?[6-9]\d{9}|\b\d{3}[-.\s]??\d{3}[-.\s]??\d{4}\b"
CGPA_REGEX = r"(?:cgpa|gpa|percentage|score)[\s:]*([0-9]+(?:\.[0-9]+)?)"
YEAR_REGEX = r"\b(20[1-3][0-9])\b"


def extract_text_from_pdf(file_bytes: bytes) -> Tuple_Text_Or_Error:
    text = ""
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        num_pages = len(reader.pages)
        if num_pages == 0:
            return "", ["PDF file contains zero pages."]

        for page in reader.pages:
            page_text = page.extract_text() or ""
            text += page_text + "\n"

        clean_text = text.strip()
        # Detect scanned PDFs (pages exist but text is empty or negligible)
        if num_pages > 0 and len(clean_text) < 30:
            return clean_text, [
                "Scanned or image-based PDF detected. Optical Character Recognition (OCR) is not configured. Please upload a text-selectable PDF or DOCX file."
            ]

        return clean_text, []
    except Exception as e:
        return "", [f"Malformed or unreadable PDF: {str(e)}"]


def extract_text_from_docx(file_bytes: bytes) -> Tuple_Text_Or_Error:
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        full_text = []
        for para in doc.paragraphs:
            if para.text:
                full_text.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        full_text.append(cell.text)
        clean_text = "\n".join(full_text).strip()
        if not clean_text:
            return "", ["Word document appears empty or could not be parsed."]
        return clean_text, []
    except Exception as e:
        return "", [f"Malformed or unreadable DOCX: {str(e)}"]


Tuple_Text_Or_Error = tuple[str, List[str]]


def parse_resume_content(file_bytes: bytes, filename: str) -> ResumeParseResponse:
    filename_lower = filename.lower()
    errors: List[str] = []
    text = ""

    if filename_lower.endswith(".pdf"):
        text, extract_errors = extract_text_from_pdf(file_bytes)
        errors.extend(extract_errors)
    elif filename_lower.endswith(".docx") or filename_lower.endswith(".doc"):
        text, extract_errors = extract_text_from_docx(file_bytes)
        errors.extend(extract_errors)
    else:
        return ResumeParseResponse(
            success=False,
            parsing_errors=[f"Unsupported file format '{filename}'. Allowed formats: PDF (.pdf) and DOCX (.docx)."]
        )

    if not text and errors:
        return ResumeParseResponse(
            success=False,
            parsing_errors=errors,
            raw_text_preview=""
        )

    # 1. Extract contact details
    email_match = re.search(EMAIL_REGEX, text)
    email = email_match.group(0) if email_match else None

    phone_match = re.search(PHONE_REGEX, text)
    phone = phone_match.group(0) if phone_match else None

    # Candidate Name heuristic (first clean line without numbers, emails or URLs)
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    candidate_name = None
    for line in lines[:6]:
        if (
            len(line.split()) <= 4
            and not re.search(r"[@0-9|:]", line)
            and "resume" not in line.lower()
            and "curriculum" not in line.lower()
        ):
            candidate_name = line
            break

    # 2. Extract Skills using taxonomy and word boundary checks
    lower_text = text.lower()
    matched_skills = set()
    for skill in KNOWN_SKILLS:
        pattern = r"\b" + re.escape(skill) + r"\b"
        if re.search(pattern, lower_text):
            matched_skills.add(normalize_token(skill))

    # 3. Extract Education details
    education_list: List[ExtractedEducation] = []
    deg_patterns = [
        r"(B\.?Tech|Bachelor of Technology|B\.?E\.?|Bachelor of Engineering|M\.?Tech|MCA|BCA|B\.?Sc|M\.?Sc)",
    ]
    extracted_degree = None
    for pattern in deg_patterns:
        deg_match = re.search(pattern, text, re.IGNORECASE)
        if deg_match:
            extracted_degree = deg_match.group(0)
            break

    cgpa_val = None
    cgpa_match = re.search(CGPA_REGEX, text, re.IGNORECASE)
    if cgpa_match:
        try:
            val = float(cgpa_match.group(1))
            if val <= 10.0:
                cgpa_val = val
            elif val <= 100.0:
                # Percentage to 10-point scale approx
                cgpa_val = round(val / 10.0, 2)
        except ValueError:
            pass

    year_val = None
    years = re.findall(YEAR_REGEX, text)
    if years:
        year_val = int(years[-1])

    if extracted_degree or cgpa_val or year_val:
        education_list.append(
            ExtractedEducation(
                degree=extracted_degree,
                institution=None,
                graduation_year=year_val,
                cgpa=cgpa_val,
            )
        )

    # 4. Extract Projects
    projects_list: List[ExtractedProject] = []
    project_sections = re.findall(
        r"(?:project|projects)[\s:]*([^\n\r]+)(?:[\n\r]+([^\n\r]+))?",
        text,
        re.IGNORECASE
    )
    for p_title, p_desc in project_sections[:3]:
        clean_title = p_title.strip()
        if len(clean_title) > 3 and len(clean_title) < 80:
            techs = [s for s in matched_skills if s in (clean_title + " " + (p_desc or "")).lower()]
            projects_list.append(
                ExtractedProject(
                    title=clean_title,
                    description=p_desc.strip() if p_desc else None,
                    technologies=techs,
                )
            )

    # 5. Certifications
    certifications: List[str] = []
    cert_matches = re.findall(
        r"(?:certified|certification|certificate)[\s:]*([^\n\r]+)",
        text,
        re.IGNORECASE
    )
    for c in cert_matches[:4]:
        c_clean = c.strip()
        if len(c_clean) > 3 and len(c_clean) < 100:
            certifications.append(c_clean)

    # 6. Experience
    experience_list: List[ExtractedExperience] = []
    exp_matches = re.findall(
        r"(?:intern|internship|software engineer|developer|analyst)[\s:]*([^\n\r]+)",
        text,
        re.IGNORECASE
    )
    for exp in exp_matches[:3]:
        exp_clean = exp.strip()
        if len(exp_clean) > 3 and len(exp_clean) < 80:
            experience_list.append(
                ExtractedExperience(
                    role=exp_clean,
                    organization=None,
                    duration=None,
                    description=None,
                )
            )

    preview_text = text[:800] + ("..." if len(text) > 800 else "")

    return ResumeParseResponse(
        success=True,
        candidate_name=candidate_name,
        email=email,
        phone=phone,
        skills=sorted(list(matched_skills)),
        education=education_list,
        projects=projects_list,
        certifications=certifications,
        experience=experience_list,
        raw_text_preview=preview_text,
        parsing_errors=errors,
    )
